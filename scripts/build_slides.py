#!/usr/bin/env python
"""Build writing/Digital-Minds-Sprint-Slides.pptx from writing/slides.md.

slides.md format — one slide per '# ' heading:

    # Slide title
    subtitle: optional one-liner under the title
    layout: title | bullets | image | image-left | two-col | table | quote
    image: writing/figures/fig2_unbuildable.png       (for image layouts)
    caption: text under the image
    - bullet
      - sub-bullet
    | a | b |   (pipe table rows for layout: table)
    notes: speaker notes (may repeat lines)

`{{KEY}}` placeholders are substituted from writing/paper_fill.json.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

ROOT = Path(__file__).resolve().parents[1]
W = ROOT / "writing"
OUT = W / "Digital-Minds-Sprint-Slides.pptx"

INK = RGBColor(0x1F, 0x1F, 0x1F)
MUTED = RGBColor(0x5B, 0x5B, 0x5B)
ACCENT = RGBColor(0x1B, 0x4F, 0x9C)
ACCENT2 = RGBColor(0xB2, 0x3A, 0x1D)
BG = RGBColor(0xFA, 0xF8, 0xF3)
FONT = "Calibri"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def parse(md: str, fill: dict):
    def sub(t):
        return re.sub(r"\{\{\s*([A-Za-z0-9_]+)\s*\}\}", lambda m: str(fill.get(m.group(1), f"[[{m.group(1)}]]")), t)
    md = sub(sub(sub(md)))
    slides = []
    cur = None
    for line in md.split("\n"):
        if line.startswith("# "):
            cur = {"title": line[2:].strip(), "bullets": [], "table": [], "notes": [], "layout": "bullets",
                   "image": None, "caption": None, "subtitle": None, "right": []}
            slides.append(cur)
            continue
        if cur is None or not line.strip():
            continue
        m = re.match(r"^(layout|image|caption|subtitle|notes|right):\s*(.*)$", line)
        if m:
            k, v = m.group(1), m.group(2).strip()
            if k == "notes":
                cur["notes"].append(v)
            elif k == "right":
                cur["right"].append(v)
            else:
                cur[k] = v
            continue
        if line.lstrip().startswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            cur["table"].append(cells)
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            depth = len(m.group(1)) // 2
            cur["bullets"].append((depth, m.group(2).strip()))
            continue
        cur["bullets"].append((0, line.strip()))
    return slides


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_runs(par, text, size, color=INK, bold=False):
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        b, i = bold, False
        if tok.startswith("**") and tok.endswith("**"):
            tok, b = tok[2:-2], True
        elif tok.startswith("*") and tok.endswith("*"):
            tok, i = tok[1:-1], True
        elif tok.startswith("`") and tok.endswith("`"):
            tok = tok[1:-1]
        r = par.add_run()
        r.text = tok
        r.font.size = Pt(size)
        r.font.bold = b
        r.font.italic = i
        r.font.name = FONT
        r.font.color.rgb = color


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def title_bar(slide, title, subtitle=None, small=False):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), SLIDE_W - Inches(1.2), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    add_runs(p, title, 30 if not small else 24, ACCENT, bold=True)
    if subtitle:
        p2 = tf.add_paragraph()
        add_runs(p2, subtitle, 16, MUTED)
    # rule
    ln = slide.shapes.add_shape(1, Inches(0.6), Inches(1.28), SLIDE_W - Inches(1.2), Emu(18000))
    ln.fill.solid()
    ln.fill.fore_color.rgb = ACCENT
    ln.line.fill.background()


def bullets_box(slide, items, left, top, width, height, size=18):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for depth, text in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = min(depth, 4)
        marker = "• " if depth == 0 else "– "
        add_runs(p, marker + text, size - 2 * depth, INK)
        p.space_after = Pt(6)
    return tb


def image_box(slide, path, left, top, max_w, max_h):
    full = ROOT / path
    if not full.exists():
        tb = slide.shapes.add_textbox(left, top, max_w, Inches(0.5))
        add_runs(tb.text_frame.paragraphs[0], f"[[MISSING {path}]]", 14, ACCENT2)
        return None
    from PIL import Image
    with Image.open(full) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    pic = slide.shapes.add_picture(str(full), left, top, width=int(w * scale), height=int(h * scale))
    return pic


def caption_box(slide, text, left, top, width):
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.6))
    tf = tb.text_frame
    tf.word_wrap = True
    add_runs(tf.paragraphs[0], text, 12, MUTED)


def table_box(slide, rows, left, top, width, height, size=12):
    n_r, n_c = len(rows), max(len(r) for r in rows)
    shape = slide.shapes.add_table(n_r, n_c, left, top, width, height)
    t = shape.table
    for i, row in enumerate(rows):
        for j in range(n_c):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            add_runs(p, row[j] if j < len(row) else "", size, INK, bold=(i == 0))
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            if i == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xDD, 0xE6, 0xF2)
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i % 2 else RGBColor(0xF3, 0xF3, 0xF3)
    return shape


def build(slides):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]
    for k, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        bg(sl)
        lay = s["layout"]
        if lay == "title":
            tb = sl.shapes.add_textbox(Inches(0.8), Inches(2.0), SLIDE_W - Inches(1.6), Inches(2.2))
            tf = tb.text_frame
            tf.word_wrap = True
            add_runs(tf.paragraphs[0], s["title"], 40, ACCENT, bold=True)
            if s["subtitle"]:
                p = tf.add_paragraph()
                add_runs(p, s["subtitle"], 20, MUTED)
            bullets_box(sl, s["bullets"], Inches(0.8), Inches(4.5), SLIDE_W - Inches(1.6), Inches(2.5), size=16)
        elif lay == "quote":
            title_bar(sl, s["title"], s["subtitle"])
            tb = sl.shapes.add_textbox(Inches(1.2), Inches(2.0), SLIDE_W - Inches(2.4), Inches(4.5))
            tf = tb.text_frame
            tf.word_wrap = True
            first = True
            for depth, text in s["bullets"]:
                p = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                add_runs(p, text, 22, INK)
                p.space_after = Pt(14)
        elif lay == "image":
            title_bar(sl, s["title"], s["subtitle"])
            top = Inches(1.5)
            avail_h = SLIDE_H - top - Inches(1.1)
            if s["bullets"]:
                bullets_box(sl, s["bullets"], Inches(0.6), top, SLIDE_W - Inches(1.2), Inches(1.0), size=15)
                top += Inches(0.25 + 0.3 * len(s["bullets"]))
                avail_h = SLIDE_H - top - Inches(1.0)
            pic = image_box(sl, s["image"], Inches(0.6), top, SLIDE_W - Inches(1.2), avail_h)
            if pic is not None:
                # centre horizontally
                pic.left = int((SLIDE_W - pic.width) / 2)
                cap_top = pic.top + pic.height + Inches(0.05)
            else:
                cap_top = top + Inches(0.6)
            if s["caption"]:
                caption_box(sl, s["caption"], Inches(0.6), cap_top, SLIDE_W - Inches(1.2))
        elif lay == "image-left":
            title_bar(sl, s["title"], s["subtitle"])
            pic = image_box(sl, s["image"], Inches(0.5), Inches(1.5), Inches(7.6), Inches(5.4))
            bullets_box(sl, s["bullets"], Inches(8.3), Inches(1.5), Inches(4.6), Inches(5.5), size=16)
            if s["caption"] and pic is not None:
                caption_box(sl, s["caption"], Inches(0.5), pic.top + pic.height + Inches(0.05), Inches(7.6))
        elif lay == "two-col":
            title_bar(sl, s["title"], s["subtitle"])
            bullets_box(sl, s["bullets"], Inches(0.6), Inches(1.5), Inches(6.0), Inches(5.6), size=17)
            right = [(0, t) for t in s["right"]]
            bullets_box(sl, right, Inches(6.9), Inches(1.5), Inches(5.9), Inches(5.6), size=17)
        elif lay == "table":
            title_bar(sl, s["title"], s["subtitle"])
            top = Inches(1.5)
            if s["bullets"]:
                bullets_box(sl, s["bullets"], Inches(0.6), top, SLIDE_W - Inches(1.2), Inches(1.0), size=15)
                top += Inches(0.3 + 0.32 * len(s["bullets"]))
            if s["table"]:
                table_box(sl, s["table"], Inches(0.6), top, SLIDE_W - Inches(1.2),
                          Inches(0.36) * len(s["table"]), size=12 if len(s["table"][0]) <= 6 else 10)
            if s["caption"]:
                caption_box(sl, s["caption"], Inches(0.6), SLIDE_H - Inches(0.9), SLIDE_W - Inches(1.2))
        else:  # bullets
            title_bar(sl, s["title"], s["subtitle"])
            bullets_box(sl, s["bullets"], Inches(0.6), Inches(1.5), SLIDE_W - Inches(1.2), Inches(5.6), size=18)
        # footer
        ft = sl.shapes.add_textbox(Inches(0.6), SLIDE_H - Inches(0.45), SLIDE_W - Inches(1.2), Inches(0.3))
        add_runs(ft.text_frame.paragraphs[0], f"Digital Minds Research Sprint · {k+1}/{len(slides)}", 10, MUTED)
        if s["notes"]:
            sl.notes_slide.notes_text_frame.text = "\n".join(s["notes"])
    prs.save(OUT)
    print("wrote", OUT, len(slides), "slides")


if __name__ == "__main__":
    import sys
    # Optional positional overrides: slides.md fill.json out.pptx (writing/v2 uses them).
    md_path = Path(sys.argv[1]) if len(sys.argv) > 1 else W / "slides.md"
    fp = Path(sys.argv[2]) if len(sys.argv) > 2 else W / "paper_fill.json"
    if len(sys.argv) > 3:
        OUT = Path(sys.argv[3])
    fill = {}
    if fp.exists():
        fill = json.loads(fp.read_text(encoding="utf-8"))
    build(parse(md_path.read_text(encoding="utf-8"), fill))
